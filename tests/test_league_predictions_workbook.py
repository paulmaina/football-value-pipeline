"""Display-only league predictions workbook written beside the readable CSVs."""
from pathlib import Path
import json
import re
import tempfile
import unittest
from unittest.mock import patch
import warnings
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import pandas as pd

from offline_test_case import OfflineQuotePolicyTestCase
import format_football_output as formatter
from format_football_output import WORKBOOK_NAME, main, write_predictions_workbook
from test_conservative_bet_filter import fixture
from test_readable_safest_ah import MARGINS, offer

try:
    import openpyxl
except ImportError:  # pragma: no cover - CI installs only the runtime requirements
    openpyxl = None

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
LEAD = ["Match date", "League", "Home team", "Away team", "Model lean"]
UNSIGNED_PCT = re.compile(r"^\d+(\.\d+)?%$")


def league_rows(league, count, start=0):
    rows = []
    for i in range(start, start + count):
        row = fixture(i)
        row.update({
            "league": league,
            "fixture_source": "The Odds API",
            "p_over_2_5": 0.55, "p_under_2_5": 0.45, "p_btts_yes": 0.52, "p_btts_no": 0.48,
            "p_over_0_5": 0.93, "p_over_3_5": 0.31, "p_under_3_5": 0.69, "p_under_4_5": 0.84,
            "p_home_draw": 0.89, "p_draw_away": 0.28, "p_home_away": 0.83,
            "lambda_home": 1.9, "lambda_away": 0.8,
            "goal_margin_distribution": MARGINS,
            "asian_handicap_offers": json.dumps([offer("home", .5, 1.40, book="Pinnacle", opposite=2.9)]),
            "totals_anchor_status": "ANCHORED_MAIN",
            "top_scorelines": json.dumps([{"score": "2-0", "p": 0.137}, {"score": "1-0", "p": 0.12}]),
        })
        rows.append(row)
    return rows


class WorkbookCase(OfflineQuotePolicyTestCase):
    """Runs the formatter on EPL and LaLiga fixtures, an empty UCL file and a stale UEL CSV."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        root = Path(self.tmp.name)
        raw, self.readable = root / "raw", root / "readable"
        raw.mkdir()
        self.readable.mkdir()
        pd.DataFrame(league_rows("Premier League", 3)).to_csv(raw / "EPL_predictions.csv", index=False)
        pd.DataFrame(league_rows("La Liga", 2, start=10)).to_csv(raw / "LaLiga_predictions.csv", index=False)
        (raw / "UCL_predictions.csv").write_text("\n", encoding="utf-8")
        # A stale readable CSV from an earlier run holds already-played matches.
        stale = pd.DataFrame(league_rows("Europa League", 1, start=20))
        stale.to_csv(raw / "stale.csv", index=False)
        formatter.readable_predictions_frame(stale, {}).to_csv(
            self.readable / "UEL_predictions_readable.csv", index=False, encoding="utf-8-sig")
        (raw / "stale.csv").unlink()
        self.assertEqual(main(["--root-dir", str(root)]), 0)
        self.path = self.readable / WORKBOOK_NAME

    def csv(self, league):
        return pd.read_csv(self.readable / f"{league}_predictions_readable.csv",
                           dtype=str, keep_default_na=False, encoding="utf-8-sig")

    def parts(self):
        with ZipFile(self.path) as archive:
            self.assertIsNone(archive.testzip())
            return {name: archive.read(name) for name in archive.namelist()}

    def sheet_xmls(self):
        parts = self.parts()
        return [ET.fromstring(parts[n]) for n in sorted(parts) if n.startswith("xl/worksheets/")]


class LeaguePredictionsWorkbookTests(WorkbookCase):
    def test_every_fixture_league_gets_a_tab_and_empty_or_stale_leagues_do_not(self):
        workbook = ET.fromstring(self.parts()["xl/workbook.xml"])
        names = [s.get("name") for s in workbook.findall("m:sheets/m:sheet", NS)]
        self.assertEqual(sorted(names), ["EPL", "LaLiga"])
        self.assertNotIn("UCL", names)
        self.assertNotIn("UEL", names)

    def test_fixture_source_is_only_in_the_csvs(self):
        for league in ("EPL", "LaLiga"):
            self.assertIn("Fixture source", self.csv(league).columns)
        for xml in self.sheet_xmls():
            headers = [c.findtext(".//m:t", namespaces=NS) for c in xml.findall("m:sheetData/m:row[@r='1']/m:c", NS)]
            self.assertNotIn("Fixture source", headers)
            self.assertEqual(headers[:5], LEAD)

    def test_frozen_pane_is_consistent_with_its_split(self):
        # Excel repairs a frozen pane whose top-left cell sits inside the frozen area.
        for xml in self.sheet_xmls():
            pane = xml.find("m:sheetViews/m:sheetView/m:pane", NS)
            self.assertEqual((pane.get("xSplit"), pane.get("ySplit")), ("5", "1"))
            self.assertEqual(pane.get("topLeftCell"), "F2")
            self.assertEqual(pane.get("state"), "frozen")
            self.assertEqual(pane.get("activePane"), "bottomRight")
            panes = {s.get("pane") for s in xml.findall("m:sheetViews/m:sheetView/m:selection", NS)}
            self.assertEqual(panes, {"topRight", "bottomLeft", "bottomRight"})
            self.assertIsNotNone(xml.find("m:autoFilter", NS))
            # Schema order: autoFilter follows sheetData.
            tags = [child.tag.split("}")[1] for child in xml]
            self.assertLess(tags.index("sheetData"), tags.index("autoFilter"))
        sheets = [x.find("m:sheetViews/m:sheetView", NS).get("tabSelected") for x in self.sheet_xmls()]
        self.assertEqual(sheets.count("1"), 1)

    def test_every_part_is_well_formed_xml(self):
        for name, data in self.parts().items():
            with self.subTest(part=name):
                ET.fromstring(data)

    def test_failed_write_keeps_the_published_workbook(self):
        before = self.path.read_bytes()
        with patch.object(formatter.ZipFile, "writestr", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                write_predictions_workbook([("EPL", self.csv("EPL"))], self.path)
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual([p.name for p in self.readable.glob("*.tmp")], [])

    def test_no_fixtures_anywhere_replaces_old_tabs_with_a_note(self):
        write_predictions_workbook([], self.path)
        workbook = ET.fromstring(self.parts()["xl/workbook.xml"])
        self.assertEqual([s.get("name") for s in workbook.findall("m:sheets/m:sheet", NS)], ["No fixtures"])


@unittest.skipIf(openpyxl is None, "openpyxl is not installed")
class LeaguePredictionsWorkbookOpenpyxlTests(WorkbookCase):
    def load(self):
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            return openpyxl.load_workbook(self.path)

    def test_reloads_cleanly_with_openpyxl(self):
        workbook = self.load()
        self.assertEqual(sorted(workbook.sheetnames), ["EPL", "LaLiga"])
        for sheet in workbook.worksheets:
            self.assertEqual(sheet.freeze_panes, "F2")
            self.assertEqual([c.value for c in sheet[1][:5]], LEAD)
            self.assertEqual(sheet.auto_filter.ref, f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}")
            self.assertTrue(all(d.width >= 10 for d in sheet.column_dimensions.values()))

    def test_rows_and_values_match_the_csv(self):
        workbook = self.load()
        for league in ("EPL", "LaLiga"):
            frame, sheet = self.csv(league), workbook[league]
            headers = [c.value for c in sheet[1]]
            self.assertEqual(sorted(headers), sorted(c for c in frame.columns if c != "Fixture source"))
            self.assertEqual(sheet.max_row - 1, len(frame))
            for r, record in enumerate(frame.to_dict("records"), start=2):
                for c, header in enumerate(headers, start=1):
                    text, cell = record[header], sheet.cell(r, c)
                    with self.subTest(league=league, column=header, row=r):
                        if text in ("", "—"):
                            self.assertIsNone(cell.value)
                        elif cell.number_format in ("0.0%", "+0.0%;-0.0%;0.0%"):
                            self.assertAlmostEqual(cell.value, float(text.rstrip("%")) / 100)
                        elif cell.number_format == "+0.0;-0.0;0.0":
                            self.assertAlmostEqual(cell.value, float(text.removesuffix(" pp")))
                        elif cell.number_format == "0.00":
                            self.assertAlmostEqual(cell.value, float(text))
                        elif cell.number_format == "yyyy-mm-dd":
                            self.assertEqual(cell.value.date().isoformat(), text)
                        else:
                            self.assertEqual(cell.value, text)

    def test_probability_gap_and_odds_columns_are_numeric(self):
        workbook = self.load()
        frame, sheet = self.csv("EPL"), workbook["EPL"]
        headers = [c.value for c in sheet[1]]
        # Any column the CSV shows as an unsigned percentage is a 0-1 probability.
        probability = [h for h in headers if (frame[h] != "—").any()
                       and frame[h][frame[h] != "—"].str.match(UNSIGNED_PCT).all()]
        for expected in ["Model lean probability", "Home win", "Draw", "Away win",
                         "Most likely score probability", "Over 0.5", "Over 1.5", "Over 2.5",
                         "Under 2.5", "Over 3.5", "Under 3.5", "Under 4.5",
                         "Double Chance 1X", "Double Chance X2", "Double Chance 12", "BTTS Yes",
                         "BTTS No", "Safest AH model no-loss probability",
                         "Safest AH model break-even win share", "Safest AH Pinnacle no-vig win share"]:
            self.assertIn(expected, probability)
        formats = {
            **dict.fromkeys(probability, "0.0%"),
            "Safest AH model minus Pinnacle win share": "+0.0;-0.0;0.0",
            "Home market odds": "0.00", "Away market odds": "0.00",
        }
        for header, number_format in formats.items():
            cell = sheet.cell(2, headers.index(header) + 1)
            with self.subTest(column=header):
                self.assertEqual(cell.data_type, "n")
                self.assertIsInstance(cell.value, float)
                self.assertEqual(cell.number_format, number_format)
                if number_format == "0.0%":
                    self.assertTrue(0 <= cell.value <= 1)


if __name__ == "__main__":
    unittest.main()
